# -*- coding: utf-8 -*-
"""Markdown 笔记 - 纯原生 PySide6 桌面应用

完全用 Qt Widgets 实现，不依赖任何浏览器内核：
  - 左侧笔记列表（标题 / 大小 / 更新时间）
  - 顶部工具栏：导出 MD / PDF / Word、删除笔记
  - 中间 Markdown 编辑器，右侧实时预览
  - 数据持久化到本地 JSON 文件

启动方式：pythonw.exe main.pyw（无控制台黑框）。
"""

import os
import sys
import json
import time
import uuid
import traceback
import re
from datetime import datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "notes_data")
DATA_FILE = os.path.join(DATA_DIR, "notes.json")
CRASH_LOG = os.path.join(BASE_DIR, "crash.log")

# 主题色
PRIMARY = "#7c3aed"          # 紫色主色
PRIMARY_HOVER = "#6d28d9"    # 深紫
PRIMARY_LIGHT = "#ede9fe"    # 浅紫背景
BG = "#ffffff"               # 主背景
SIDEBAR_BG = "#f8fafc"       # 侧边栏背景
BORDER = "#e2e8f0"           # 边框
TEXT = "#1e293b"             # 主文字
TEXT_SECONDARY = "#64748b"   # 次要文字
TEXT_MUTED = "#94a3b8"       # 更淡文字


def show_error(title: str, text: str) -> None:
    try:
        import ctypes
        ctypes.windll.user32.MessageBoxW(0, str(text), str(title), 0x10)
    except Exception:
        pass


def handle_exception(exc_type, exc_value, exc_tb):
    msg = "".join(traceback.format_exception(exc_type, exc_value, exc_tb))
    try:
        with open(CRASH_LOG, "a", encoding="utf-8") as f:
            f.write("\n=== %s ===\n" % time.strftime("%Y-%m-%d %H:%M:%S"))
            f.write(msg)
    except Exception:
        pass
    show_error("Markdown 笔记 - 错误", msg)


sys.excepthook = handle_exception


# ---------------------------------------------------------------------------
# 数据层
# ---------------------------------------------------------------------------
def load_notes():
    if not os.path.isfile(DATA_FILE):
        return []
    try:
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return data
    except Exception:
        pass
    return []


def save_notes(notes):
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(notes, f, ensure_ascii=False, indent=2)


def derive_title(content: str) -> str:
    for line in content.splitlines():
        line = line.strip().lstrip("#").strip()
        if line:
            return line[:40]
    return "未命名笔记"


def format_time(ts: float) -> str:
    if not ts:
        return ""
    try:
        dt = datetime.fromtimestamp(ts)
        now = datetime.now()
        if dt.date() == now.date():
            return dt.strftime("今天 %H:%M")
        elif (now.date() - dt.date()).days == 1:
            return "昨天"
        elif now.year == dt.year:
            return dt.strftime("%m月%d日")
        else:
            return dt.strftime("%Y年%m月%d日")
    except Exception:
        return ""


def format_size(content: str) -> str:
    size = len(content.encode("utf-8"))
    if size < 1024:
        return "%d B" % size
    return "%.1f KB" % (size / 1024)


def word_count(content: str) -> int:
    # 中文按字，英文按词
    cn = len(re.findall(r"[\u4e00-\u9fff]", content))
    en = len(re.findall(r"[a-zA-Z0-9_]+", content))
    return cn + en


# ---------------------------------------------------------------------------
# 导出
# ---------------------------------------------------------------------------
def export_md(note, save_path: str):
    with open(save_path, "w", encoding="utf-8") as f:
        f.write(note.get("content", ""))


def export_pdf(note, save_path: str):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.enums import TA_LEFT

    # 尝试注册中文字体；失败则使用默认 Helvetica（英文可显示，中文可能乱码）
    font_name = "Helvetica"
    possible_fonts = [
        ("C:/Windows/Fonts/msyh.ttc", "MicrosoftYaHei"),
        ("C:/Windows/Fonts/simhei.ttf", "SimHei"),
        ("C:/Windows/Fonts/simsun.ttc", "SimSun"),
    ]
    for path, name in possible_fonts:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                font_name = name
                break
            except Exception:
                pass

    doc = SimpleDocTemplate(save_path, pagesize=A4,
                            rightMargin=2 * cm, leftMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="CN", fontName=font_name, fontSize=11, leading=18,
        alignment=TA_LEFT, spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name="CN_H1", fontName=font_name, fontSize=20, leading=28,
        textColor=PRIMARY, spaceAfter=12,
    ))
    styles.add(ParagraphStyle(
        name="CN_H2", fontName=font_name, fontSize=16, leading=22,
        textColor=PRIMARY, spaceAfter=10,
    ))
    styles.add(ParagraphStyle(
        name="CN_H3", fontName=font_name, fontSize=14, leading=20,
        textColor=PRIMARY, spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="CN_CODE", fontName="Courier", fontSize=9, leading=12,
        backColor="#f1f5f9", leftIndent=6, rightIndent=6,
        spaceAfter=6, borderPadding=6,
    ))

    story = []
    content = note.get("content", "")

    # 按行解析 markdown，简单分段渲染
    in_code = False
    code_lines = []
    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code:
                story.append(Preformatted("\n".join(code_lines), styles["CN_CODE"]))
                story.append(Spacer(1, 6))
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            story.append(Spacer(1, 6))
            continue

        # 标题
        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["CN_H1"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["CN_H2"]))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], styles["CN_H3"]))
        elif line.startswith("> "):
            story.append(Paragraph("<i>%s</i>" % line[2:], styles["CN"]))
        elif line.startswith("- ") or line.startswith("* "):
            story.append(Paragraph("• %s" % line[2:], styles["CN"]))
        elif re.match(r"^\d+\.\s", line):
            story.append(Paragraph(line, styles["CN"]))
        else:
            # 粗体 / 斜体简单替换
            html = line
            html = re.sub(r"\*\*\*(.+?)\*\*\*", r"<b><i>\1</i></b>", html)
            html = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html)
            html = re.sub(r"\*(.+?)\*", r"<i>\1</i>", html)
            html = re.sub(r"`(.+?)`", r"<font face='Courier' backColor='#f1f5f9'>\1</font>", html)
            story.append(Paragraph(html, styles["CN"]))

    if in_code and code_lines:
        story.append(Preformatted("\n".join(code_lines), styles["CN_CODE"]))

    doc.build(story)


def export_docx(note, save_path: str):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    doc = Document()

    # 设置默认中文字体
    try:
        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set("w:eastAsia", "Microsoft YaHei")
        style.font.size = Pt(11)
    except Exception:
        pass

    content = note.get("content", "")
    in_code = False
    code_lines = []

    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Inches(0.2)
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=1)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x7c, 0x3a, 0xed)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], level=2)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x7c, 0x3a, 0xed)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=3)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x7c, 0x3a, 0xed)
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.paragraph_format.left_indent = Inches(0.2)
            for run in p.runs:
                run.italic = True
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
        else:
            p = doc.add_paragraph()
            # 简单解析粗体/斜体/行内代码
            parts = re.split(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", line)
            for part in parts:
                run = p.add_run(part)
                if part.startswith("***") and part.endswith("***"):
                    run.text = part[3:-3]
                    run.bold = True
                    run.italic = True
                elif part.startswith("**") and part.endswith("**"):
                    run.text = part[2:-2]
                    run.bold = True
                elif part.startswith("*") and part.endswith("*"):
                    run.text = part[1:-1]
                    run.italic = True
                elif part.startswith("`") and part.endswith("`"):
                    run.text = part[1:-1]
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)

    if in_code and code_lines:
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Courier New"
        run.font.size = Pt(9)

    doc.save(save_path)


# ---------------------------------------------------------------------------
# 导入：PDF / Word → Markdown
# ---------------------------------------------------------------------------
def import_pdf_to_md(path: str) -> str:
    from pdfminer.high_level import extract_text

    raw = extract_text(path) or ""
    lines = [ln.rstrip() for ln in raw.splitlines()]
    out = []
    blank = 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank <= 1:
                out.append("")
            continue
        blank = 0
        out.append(ln.strip())
    text = "\n".join(out).strip()
    return text


def _docx_run_to_md(run) -> str:
    text = run.text
    if not text:
        return ""
    if run.bold and run.italic:
        text = "***%s***" % text
    elif run.bold:
        text = "**%s**" % text
    elif run.italic:
        text = "*%s*" % text
    if run.underline:
        text = "<u>%s</u>" % text
    return text


def import_docx_to_md(path: str) -> str:
    from docx import Document

    doc = Document(path)
    out = []

    def _style_level(p):
        # 返回标题级别 1-6，否则 0
        try:
            name = (p.style.name or "") if p.style else ""
        except Exception:
            name = ""
        if name.startswith("Heading ") or name.startswith("标题 "):
            try:
                return int(name.split()[-1])
            except Exception:
                return 0
        if name in ("Title", "Subtitle"):
            return 1
        return 0

    for block in doc.element.body.iterchildren():
        from docx.oxml.ns import qn
        if block.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            p = Paragraph(block, doc)
            level = _style_level(p)
            prefix = ""
            # 列表检测
            style_name = ""
            try:
                style_name = (p.style.name or "") if p.style else ""
            except Exception:
                pass
            is_bullet = style_name.lower().startswith("list bullet") or "列表" in style_name
            is_number = style_name.lower().startswith("list number") or style_name.lower().startswith("list decimal")

            if level > 0:
                prefix = "#" * min(level, 6) + " "
            elif is_bullet:
                prefix = "- "
            elif is_number:
                prefix = "1. "

            md_text = "".join(_docx_run_to_md(r) for r in p.runs)
            if not md_text and p.text:
                md_text = p.text
            if prefix or md_text.strip():
                out.append(prefix + md_text)
        elif block.tag == qn("w:tbl"):
            from docx.table import Table
            tbl = Table(block, doc)
            rows = list(tbl.rows)
            if not rows:
                continue
            header = [" ".join(c.text for c in rows[0].cells)]
            max_cols = max(len(r.cells) for r in rows)
            header_cells = [c.text.strip() for c in rows[0].cells]
            out.append("| " + " | ".join(header_cells) + " |")
            out.append("| " + " | ".join(["---"] * max(len(header_cells), 1)) + " |")
            for r in rows[1:]:
                cells = [c.text.strip() for c in r.cells]
                while len(cells) < len(header_cells):
                    cells.append("")
                out.append("| " + " | ".join(cells) + " |")
            out.append("")

    return "\n".join(out).strip()


def _read_text_file(path: str) -> str:
    """以最健壮的方式读取文本文件，自动尝试多种常见编码。"""
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030", "latin-1"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode("utf-8", errors="replace")


# ---------------------------------------------------------------------------
# 主窗口
# ---------------------------------------------------------------------------
def build_app():
    from PySide6.QtCore import Qt, QTimer, QSize
    from PySide6.QtGui import QIcon
    from PySide6.QtWidgets import (
        QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
        QListWidget, QListWidgetItem, QLineEdit, QPlainTextEdit,
        QTextBrowser, QPushButton, QSplitter, QMessageBox, QLabel,
        QFileDialog, QFrame, QSizePolicy, QSpacerItem, QStatusBar,
        QAbstractItemView,
    )
    import markdown

    md = markdown.Markdown(extensions=["fenced_code", "tables", "nl2br"])

    PREVIEW_CSS = """
    <style>
      body { font-family: -apple-system, "Segoe UI", "Microsoft YaHei", sans-serif;
             font-size: 14px; line-height: 1.7; color: #1e293b; margin: 16px; }
      h1, h2, h3, h4 { margin: 0.7em 0 0.4em; line-height: 1.25; font-weight: 600; }
      h1 { font-size: 1.8em; color: #7c3aed; border-bottom: 1px solid #e2e8f0; padding-bottom: 6px; }
      h2 { font-size: 1.45em; color: #7c3aed; border-bottom: 1px solid #e2e8f0; padding-bottom: 4px; }
      h3 { font-size: 1.2em; color: #7c3aed; }
      code { background: #f3f0ff; color: #5b21b6; padding: 1px 5px; border-radius: 4px;
             font-family: Consolas, "Courier New", monospace; font-size: 0.9em; }
      pre { background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px;
            padding: 12px; overflow: auto; }
      pre code { background: none; color: #1e293b; padding: 0; }
      blockquote { border-left: 4px solid #7c3aed; margin: 0.6em 0;
                   padding: 4px 14px; color: #475569; background: #f8fafc; }
      table { border-collapse: collapse; margin: 0.8em 0; }
      th, td { border: 1px solid #e2e8f0; padding: 8px 12px; }
      th { background: #f3f0ff; color: #5b21b6; }
      a { color: #7c3aed; }
      hr { border: none; border-top: 1px solid #e2e8f0; }
      img { max-width: 100%; }
      ul, ol { padding-left: 1.4em; }
    </style>
    """

    class NoteListItem(QWidget):
        """自定义笔记列表项：标题 + 大小 + 时间"""
        def __init__(self, title, size_text, time_text, parent=None):
            super().__init__(parent)
            self.setFixedHeight(58)
            layout = QVBoxLayout(self)
            layout.setContentsMargins(12, 6, 12, 6)
            layout.setSpacing(2)

            top = QHBoxLayout()
            top.setSpacing(0)
            self.title_lbl = QLabel(title)
            self.title_lbl.setStyleSheet("font-size:14px; color:%s; font-weight:500;" % TEXT)
            self.title_lbl.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
            top.addWidget(self.title_lbl)
            layout.addLayout(top)

            bottom = QHBoxLayout()
            bottom.setSpacing(0)
            self.size_lbl = QLabel(size_text)
            self.size_lbl.setStyleSheet("font-size:11px; color:%s;" % TEXT_MUTED)
            bottom.addWidget(self.size_lbl)
            bottom.addSpacerItem(QSpacerItem(12, 1, QSizePolicy.Fixed, QSizePolicy.Minimum))
            self.time_lbl = QLabel(time_text)
            self.time_lbl.setStyleSheet("font-size:11px; color:%s;" % TEXT_MUTED)
            bottom.addWidget(self.time_lbl)
            bottom.addStretch()
            layout.addLayout(bottom)

    class MainWindow(QMainWindow):
        def __init__(self):
            super().__init__()
            self.setWindowTitle("煜辰的 Markdown 笔记")
            self.resize(1350, 860)
            self.setMinimumSize(900, 600)

            self.notes = load_notes()
            if not self.notes:
                self._seed()

            self.active_id = self.notes[0]["id"] if self.notes else None
            self._loading = False
            self._save_timer = QTimer(self)
            self._save_timer.setSingleShot(True)
            self._save_timer.timeout.connect(self._commit)

            self.setStyleSheet(self._global_stylesheet())

            central = QWidget()
            self.setCentralWidget(central)
            root = QHBoxLayout(central)
            root.setContentsMargins(0, 0, 0, 0)
            root.setSpacing(0)

            # ---------------- 左侧边栏 ----------------
            sidebar = QWidget()
            sidebar.setObjectName("sidebar")
            sidebar.setFixedWidth(280)
            sidebar_v = QVBoxLayout(sidebar)
            sidebar_v.setContentsMargins(0, 0, 0, 0)
            sidebar_v.setSpacing(0)

            # 顶部标题区
            header = QWidget()
            header.setStyleSheet("background:%s;" % SIDEBAR_BG)
            header_v = QVBoxLayout(header)
            header_v.setContentsMargins(18, 18, 18, 12)
            header_v.setSpacing(2)
            brand = QLabel("MARKDOWN")
            brand.setStyleSheet("font-size:18px; font-weight:700; color:%s; letter-spacing:1px;" % PRIMARY)
            header_v.addWidget(brand)
            subtitle = QLabel("煜辰的 Markdown")
            subtitle.setStyleSheet("font-size:12px; color:%s;" % TEXT_MUTED)
            header_v.addWidget(subtitle)
            sidebar_v.addWidget(header)

            # 搜索框
            search_wrap = QWidget()
            search_wrap.setStyleSheet("background:%s;" % SIDEBAR_BG)
            search_l = QVBoxLayout(search_wrap)
            search_l.setContentsMargins(14, 8, 14, 12)
            self.search = QLineEdit()
            self.search.setObjectName("searchBox")
            self.search.setPlaceholderText("搜索笔记…")
            self.search.textChanged.connect(self._apply_filter)
            search_l.addWidget(self.search)
            sidebar_v.addWidget(search_wrap)

            # 笔记列表
            self.list = QListWidget()
            self.list.setObjectName("noteList")
            self.list.setFocusPolicy(Qt.NoFocus)
            self.list.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
            self.list.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
            self.list.itemClicked.connect(self._on_item_clicked)
            sidebar_v.addWidget(self.list, 1)

            # 新建按钮
            btn_wrap = QWidget()
            btn_wrap.setStyleSheet("background:%s; border-top:1px solid %s;" % (SIDEBAR_BG, BORDER))
            btn_l = QVBoxLayout(btn_wrap)
            btn_l.setContentsMargins(14, 14, 14, 18)
            self.btn_new = QPushButton("+ 新建笔记")
            self.btn_new.setObjectName("btnNew")
            self.btn_new.setCursor(Qt.PointingHandCursor)
            self.btn_new.setMinimumSize(220, 44)
            self.btn_new.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            self.btn_new.setStyleSheet("""
                QPushButton {
                    background-color: #7c3aed; color: #ffffff;
                    border: 1px solid #7c3aed; border-radius: 8px;
                    min-height: 44px; padding: 0 16px;
                    font-size: 14px; font-weight: 600;
                }
                QPushButton:hover { background-color: #6d28d9; border-color: #6d28d9; }
                QPushButton:pressed { background-color: #5b21b6; border-color: #5b21b6; }
            """)
            self.btn_new.clicked.connect(self._create)
            btn_l.addWidget(self.btn_new)
            sidebar_v.addWidget(btn_wrap)

            # ---------------- 右侧主区域 ----------------
            main = QWidget()
            main_v = QVBoxLayout(main)
            main_v.setContentsMargins(0, 0, 0, 0)
            main_v.setSpacing(0)

            # 顶部工具栏
            toolbar = QWidget()
            toolbar.setObjectName("toolbar")
            toolbar_h = QHBoxLayout(toolbar)
            toolbar_h.setContentsMargins(20, 12, 20, 12)
            toolbar_h.setSpacing(0)

            self.title_lbl = QLabel("Vibe Coding 学习笔记")
            self.title_lbl.setStyleSheet("font-size:18px; font-weight:700; color:%s;" % TEXT)
            toolbar_h.addWidget(self.title_lbl)
            toolbar_h.addStretch()

            self.btn_import = QPushButton("导入")
            self.btn_export_md = QPushButton("导出 MD")
            self.btn_export_pdf = QPushButton("导出 PDF")
            self.btn_export_docx = QPushButton("导出 Word")
            self.btn_delete = QPushButton("删除")
            for btn in (self.btn_import, self.btn_export_md, self.btn_export_pdf, self.btn_export_docx, self.btn_delete):
                btn.setObjectName("toolBtn")
                btn.setCursor(Qt.PointingHandCursor)
                btn.setFixedHeight(32)
            self.btn_delete.setObjectName("toolBtnDelete")

            self.btn_import.clicked.connect(self._import)
            self.btn_export_md.clicked.connect(lambda: self._export("md"))
            self.btn_export_pdf.clicked.connect(lambda: self._export("pdf"))
            self.btn_export_docx.clicked.connect(lambda: self._export("docx"))
            self.btn_delete.clicked.connect(self._delete)

            toolbar_h.addWidget(self.btn_import)
            toolbar_h.addWidget(self.btn_export_md)
            toolbar_h.addWidget(self.btn_export_pdf)
            toolbar_h.addWidget(self.btn_export_docx)
            toolbar_h.addWidget(self.btn_delete)
            main_v.addWidget(toolbar)

            # 编辑/预览分割区
            splitter = QSplitter(Qt.Horizontal)
            splitter.setHandleWidth(1)
            splitter.setStyleSheet("QSplitter::handle { background:%s; }" % BORDER)

            # 编辑器容器（带标题）
            editor_wrap = QWidget()
            editor_v = QVBoxLayout(editor_wrap)
            editor_v.setContentsMargins(0, 0, 0, 0)
            editor_v.setSpacing(0)
            editor_header = QLabel("Markdown")
            editor_header.setFixedHeight(36)
            editor_header.setStyleSheet("background:#fff; color:%s; font-size:12px; font-weight:600; padding-left:16px; border-bottom:1px solid %s;" % (TEXT_MUTED, BORDER))
            editor_v.addWidget(editor_header)
            self.editor = QPlainTextEdit()
            self.editor.setObjectName("editor")
            self.editor.setPlaceholderText("在此输入 Markdown…")
            self.editor.textChanged.connect(self._on_edit)
            editor_v.addWidget(self.editor)
            splitter.addWidget(editor_wrap)

            # 预览容器
            preview_wrap = QWidget()
            preview_v = QVBoxLayout(preview_wrap)
            preview_v.setContentsMargins(0, 0, 0, 0)
            preview_v.setSpacing(0)
            preview_header = QLabel("预览")
            preview_header.setFixedHeight(36)
            preview_header.setStyleSheet("background:#fff; color:%s; font-size:12px; font-weight:600; padding-left:16px; border-bottom:1px solid %s;" % (TEXT_MUTED, BORDER))
            preview_v.addWidget(preview_header)
            self.preview = QTextBrowser()
            self.preview.setObjectName("preview")
            self.preview.setOpenExternalLinks(True)
            preview_v.addWidget(self.preview)
            splitter.addWidget(preview_wrap)

            splitter.setStretchFactor(0, 1)
            splitter.setStretchFactor(1, 1)
            main_v.addWidget(splitter, 1)

            # 底部状态栏
            self.statusBar().setObjectName("statusBar")
            self.statusBar().setStyleSheet("QStatusBar { background:%s; color:%s; border-top:1px solid %s; padding:4px 12px; font-size:12px; }" % (SIDEBAR_BG, TEXT_MUTED, BORDER))
            self.status_left = QLabel("行 1，列 1")
            self.status_right = QLabel("字数：0 · 已保存")
            self.statusBar().addWidget(self.status_left)
            self.statusBar().addPermanentWidget(self.status_right)

            root.addWidget(sidebar)
            root.addWidget(main, 1)

            self._fill_list()
            self._select(self.active_id)

        # ---------------- 样式 ----------------
        def _global_stylesheet(self):
            return """
            QMainWindow { background: %s; }
            QWidget#sidebar { background: %s; border-right: 1px solid %s; }
            QWidget#toolbar { background: %s; border-bottom: 1px solid %s; }
            QLineEdit#searchBox {
                background: #ffffff; border: 1px solid %s; border-radius: 8px;
                padding: 8px 12px; font-size: 13px; color: %s;
            }
            QLineEdit#searchBox:focus { border: 1px solid %s; }
            QLineEdit#searchBox::placeholder { color: %s; }
            QListWidget#noteList {
                background: %s; border: none; outline: none;
                padding: 6px; spacing: 4px;
            }
            QListWidget#noteList::item {
                background: #ffffff; border-radius: 8px;
                margin: 2px 8px 2px 8px;
            }
            QListWidget#noteList::item:selected {
                background: %s; border-left: 3px solid %s;
            }
            QListWidget#noteList::item:hover:!selected { background: #f1f5f9; }
            QPlainTextEdit#editor {
                background: #ffffff; border: none;
                padding: 14px; font-size: 14px; line-height: 1.6;
                color: %s; selection-background-color: %s;
            }
            QTextBrowser#preview {
                background: #ffffff; border: none;
                padding: 0px; color: %s;
            }
            QPushButton#toolBtn {
                background: #ffffff; color: %s; border: 1px solid %s;
                border-radius: 6px; padding: 0 14px; font-size: 13px;
                margin-left: 8px;
            }
            QPushButton#toolBtn:hover { background: %s; border-color: %s; color: %s; }
            QPushButton#toolBtnDelete {
                background: #ffffff; color: #dc2626; border: 1px solid #fecaca;
                border-radius: 6px; padding: 0 14px; font-size: 13px;
                margin-left: 8px;
            }
            QPushButton#toolBtnDelete:hover { background: #fee2e2; }
            QScrollBar:vertical {
                background: transparent; width: 8px; margin: 0px;
            }
            QScrollBar::handle:vertical {
                background: #cbd5e1; border-radius: 4px; min-height: 30px;
            }
            QScrollBar::handle:vertical:hover { background: #94a3b8; }
            """ % (
                BG, SIDEBAR_BG, BORDER,
                BG, BORDER,
                BORDER, TEXT, PRIMARY, TEXT_MUTED,
                SIDEBAR_BG, PRIMARY_LIGHT, PRIMARY,
                TEXT, PRIMARY_LIGHT,
                TEXT,
                TEXT_SECONDARY, BORDER, PRIMARY_LIGHT, PRIMARY, PRIMARY,
            )

        # ---------------- 数据 ----------------
        def _seed(self):
            welcome = (
                "# 欢迎使用 Markdown 笔记\n\n"
                "这是一个**纯原生 PySide6** 桌面应用，支持 Markdown 编辑与实时预览。\n\n"
                "## 功能\n\n"
                "- 左侧列表管理所有笔记\n"
                "- 顶部搜索框按**标题**过滤\n"
                "- 中间编辑，右侧**实时预览**\n"
                "- 顶部可导出 **MD / PDF / Word**\n"
                "- 数据自动保存到本地 `notes_data/notes.json`\n\n"
                "```python\n"
                "def hello():\n"
                "    print('hello, markdown')\n"
                "```\n"
            )
            self.notes = [{
                "id": uuid.uuid4().hex,
                "title": "欢迎使用 Markdown 笔记",
                "content": welcome,
                "updated_at": time.time(),
            }]
            save_notes(self.notes)

        # ---------------- 列表 ----------------
        def _fill_list(self):
            self._loading = True
            self.list.clear()
            q = self.search.text().strip().lower()
            for n in self.notes:
                if q and q not in n.get("title", "").lower():
                    continue
                item = QListWidgetItem()
                item.setData(Qt.UserRole, n["id"])
                item.setSizeHint(QSize(260, 58))
                widget = NoteListItem(
                    n.get("title") or "未命名笔记",
                    format_size(n.get("content", "")),
                    format_time(n.get("updated_at", 0)),
                )
                self.list.addItem(item)
                self.list.setItemWidget(item, widget)
            self._loading = False

        def _apply_filter(self):
            self._fill_list()
            self._highlight_active()

        def _highlight_active(self):
            for i in range(self.list.count()):
                item = self.list.item(i)
                widget = self.list.itemWidget(item)
                selected = item.data(Qt.UserRole) == self.active_id
                if selected:
                    self.list.setCurrentRow(i)
                if widget:
                    if selected:
                        widget.title_lbl.setStyleSheet("font-size:14px; color:%s; font-weight:600;" % PRIMARY)
                        widget.size_lbl.setStyleSheet("font-size:11px; color:%s;" % PRIMARY)
                        widget.time_lbl.setStyleSheet("font-size:11px; color:%s;" % PRIMARY)
                    else:
                        widget.title_lbl.setStyleSheet("font-size:14px; color:%s; font-weight:500;" % TEXT)
                        widget.size_lbl.setStyleSheet("font-size:11px; color:%s;" % TEXT_MUTED)
                        widget.time_lbl.setStyleSheet("font-size:11px; color:%s;" % TEXT_MUTED)

        def _on_item_clicked(self, item):
            nid = item.data(Qt.UserRole)
            if nid != self.active_id:
                self._select(nid)

        def _select(self, nid):
            note = next((n for n in self.notes if n["id"] == nid), None)
            if not note:
                self.active_id = None
                self._loading = True
                self.editor.clear()
                self._loading = False
                self.preview.setHtml(PREVIEW_CSS + "<body></body>")
                self.title_lbl.setText("")
                self.status_right.setText("没有笔记")
                return
            self.active_id = nid
            self._loading = True
            self.editor.setPlainText(note["content"])
            self._loading = False
            self._render_preview(note["content"])
            self.title_lbl.setText(note.get("title") or "未命名笔记")
            self._highlight_active()
            self._update_status()

        # ---------------- 编辑 / 预览 ----------------
        def _on_edit(self):
            if self._loading or not self.active_id:
                return
            text = self.editor.toPlainText()
            note = next((n for n in self.notes if n["id"] == self.active_id), None)
            if not note:
                return
            note["content"] = text
            note["title"] = derive_title(text)
            note["updated_at"] = time.time()
            self.title_lbl.setText(note["title"] or "未命名笔记")
            self.status_right.setText("字数：%d · 编辑中…" % word_count(text))
            # 更新列表
            for i in range(self.list.count()):
                item = self.list.item(i)
                if item.data(Qt.UserRole) == self.active_id:
                    widget = self.list.itemWidget(item)
                    if widget:
                        widget.title_lbl.setText(note["title"] or "未命名笔记")
                        widget.size_lbl.setText(format_size(text))
                        widget.time_lbl.setText(format_time(note["updated_at"]))
                    break
            self._render_preview(text)
            self._update_status()
            self._save_timer.start(400)

        def _render_preview(self, text):
            md.reset()
            body = md.convert(text)
            self.preview.setHtml(PREVIEW_CSS + "<body>" + body + "</body>")

        def _commit(self):
            save_notes(self.notes)
            if self.active_id:
                self.status_right.setText("字数：%d · 已保存" % word_count(self.editor.toPlainText()))

        def _update_status(self):
            cursor = self.editor.textCursor()
            line = cursor.blockNumber() + 1
            col = cursor.columnNumber() + 1
            self.status_left.setText("行 %d，列 %d" % (line, col))

        # ---------------- 操作 ----------------
        def _create(self):
            note = {
                "id": uuid.uuid4().hex,
                "title": "未命名笔记",
                "content": "# 未命名笔记\n\n开始写点什么…",
                "updated_at": time.time(),
            }
            self.notes.insert(0, note)
            save_notes(self.notes)
            self.search.clear()
            self._fill_list()
            self._select(note["id"])
            self.editor.setFocus()
            self.editor.selectAll()

        def _delete(self):
            if not self.active_id:
                return
            note = next((n for n in self.notes if n["id"] == self.active_id), None)
            if not note:
                return
            ans = QMessageBox.question(
                self, "删除笔记",
                '确定删除笔记「%s」吗？此操作不可撤销。' % (note["title"] or "未命名笔记"),
                QMessageBox.Yes | QMessageBox.No,
            )
            if ans != QMessageBox.Yes:
                return
            self.notes = [n for n in self.notes if n["id"] != self.active_id]
            save_notes(self.notes)
            self.active_id = self.notes[0]["id"] if self.notes else None
            self._fill_list()
            self._select(self.active_id)

        def _export(self, fmt: str):
            if not self.active_id:
                return
            note = next((n for n in self.notes if n["id"] == self.active_id), None)
            if not note:
                return

            title = re.sub(r'[\\/:*?"<>|]', "_", note.get("title") or "未命名笔记")
            filters = {
                "md": ("Markdown 文件 (*.md)", ".md"),
                "pdf": ("PDF 文件 (*.pdf)", ".pdf"),
                "docx": ("Word 文档 (*.docx)", ".docx"),
            }
            filt, ext = filters[fmt]
            default = title + ext
            path, _ = QFileDialog.getSaveFileName(
                self, "导出 %s" % fmt.upper(), default, filt
            )
            if not path:
                return
            if not path.lower().endswith(ext):
                path += ext

            try:
                if fmt == "md":
                    export_md(note, path)
                elif fmt == "pdf":
                    export_pdf(note, path)
                elif fmt == "docx":
                    export_docx(note, path)
                self.status_right.setText("已导出：%s" % os.path.basename(path))
            except Exception as e:
                QMessageBox.critical(self, "导出失败", "导出 %s 时出错：\n%s" % (fmt.upper(), str(e)))

        def _import(self):
            paths, _ = QFileDialog.getOpenFileNames(
                self, "导入文件", "",
                "Markdown / PDF / Word (*.md *.markdown *.pdf *.docx);;"
                "Markdown (*.md *.markdown);;PDF 文件 (*.pdf);;Word 文档 (*.docx)",
            )
            if not paths:
                return

            imported = []
            for path in paths:
                ext = os.path.splitext(path)[1].lower()
                try:
                    if ext in (".md", ".markdown"):
                        # Markdown 文件已带完整结构，直接保留原文；标题从首个标题/首行推导
                        md_text = _read_text_file(path)
                        base = os.path.splitext(os.path.basename(path))[0]
                        title = derive_title(md_text) or base or "导入的笔记"
                        content = md_text
                    elif ext == ".pdf":
                        md_text = import_pdf_to_md(path)
                        base = os.path.splitext(os.path.basename(path))[0]
                        title = base or "导入的笔记"
                        content = "# %s\n\n%s" % (base, md_text)
                    elif ext == ".docx":
                        md_text = import_docx_to_md(path)
                        base = os.path.splitext(os.path.basename(path))[0]
                        title = base or "导入的笔记"
                        content = "# %s\n\n%s" % (base, md_text)
                    else:
                        QMessageBox.warning(
                            self, "不支持的格式",
                            "已跳过不支持的文件：\n%s" % os.path.basename(path),
                        )
                        continue
                except Exception as e:
                    QMessageBox.critical(
                        self, "导入失败",
                        "读取文件时出错：\n%s\n\n%s" % (os.path.basename(path), str(e)),
                    )
                    continue

                md_text = (md_text or "").strip()
                if not md_text:
                    QMessageBox.warning(
                        self, "内容为空",
                        "未能从以下文件提取到文本内容：\n%s" % os.path.basename(path),
                    )
                    continue

                note = {
                    "id": uuid.uuid4().hex,
                    "title": title,
                    "content": content,
                    "updated_at": time.time(),
                }
                self.notes.insert(0, note)
                imported.append(note)

            if not imported:
                return
            save_notes(self.notes)
            self.search.clear()
            self._fill_list()
            self._select(imported[-1]["id"])
            self.editor.setFocus()
            self.status_right.setText("已导入 %d 个文件" % len(imported))

    app = QApplication(sys.argv)
    app.setApplicationName("煜辰的 Markdown 笔记")
    icon_path = os.path.join(BASE_DIR, "app_icon.ico")
    if os.path.isfile(icon_path):
        app.setWindowIcon(QIcon(icon_path))
    win = MainWindow()
    win.show()
    sys.exit(app.exec())


def main():
    try:
        build_app()
    except Exception:
        handle_exception(*sys.exc_info())


if __name__ == "__main__":
    main()
